import openpyxl
from openpyxl.chart import BarChart, Reference

wb = openpyxl.load_workbook("cleaned_sales_data.xlsx")
sheet = wb.active

chart = BarChart()
chart.type = "col"
chart.style = 10
chart.title = "Sales Report Chart"
chart.y_axis.title = "Total Amount"
chart.x_axis.title = "Items"

data = Reference(
    sheet, min_col=2, min_row=1, max_col=2, max_row=sheet.max_row
)
cats = Reference(sheet, min_col=1, min_row=2, max_row=sheet.max_row)

chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)

sheet.add_chart(chart, "E2")

wb.save("sales_report_with_chart.xlsx")

print("Excel Chart ကို အလိုအလျောက် ရေးဆွဲပြီးပါပြီ")