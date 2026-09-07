import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

wb = openpyxl.load_workbook("Calculated_sales.xlsx")
sheet = wb["SalesData"]

header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(
    start_color="1F4E78", end_color="1F4E78", fill_type="solid"
)
center_align = Alignment(horizontal="center", vertical="center")

thin_side = Side(border_style="thin", color="000000")
cell_border = Border(
    left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
)
for cell in sheet[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align

for row in sheet.iter_rows(
    min_row=1, max_row=sheet.max_row, min_col=1, max_col=4
):
    for cell in row:
        cell.border = cell_border

wb.save("styled_sales.xlsx")

print("Formatting အောင်မြင်စွာ ပြုလုပ်ပြီးပါပြီခင်ဗျာ!")