import os
import openpyxl

master_wb = openpyxl.Workbook()
master_sheet = master_wb.active
master_sheet.title = "MergedData"

master_sheet.append(["Item", "Total"])

files_to_merge = ["january_sales.xlsx", "february_sales.xlsx"]

for file_name in files_to_merge:
    if os.path.exists(file_name):
        wb = openpyxl.load_workbook(file_name)
        sheet = wb.active

        for row in sheet.iter_rows(
            min_col=2, max_row=sheet.max_row, values_only=True
        ):
            master_sheet.append(row)

master_wb.save("master_sales_summary.xlsx")

print("Excel finish merge")