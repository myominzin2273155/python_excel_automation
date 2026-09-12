import os
from bs4 import BeautifulSoup
from deep_translator import GoogleTranslator
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import requests
import time

# Translator နှင့် Excel ပြင်ဆင်ခြင်း
translator = GoogleTranslator(source="auto", target="my")
wb = Workbook()
ws = wb.active
ws.title = "All Pages Quotes"

# Header ခေါင်းစဉ်များ ထည့်သွင်းခြင်း
headers = ["ID", "Quote Text (Eng)", "Quote Text (MM)", "Author", "Tags"]
ws.append(headers)

base_url = "https://quotes.toscrape.com"
current_url = base_url
quote_id = 1
page_num = 1

try:
    while current_url:
        print(f"Fetching Page {page_num}: {current_url} ...")
        response = requests.get(current_url, timeout=10)

        if response.status_code != 200:
            print(f"Failed to fetch page {page_num}")
            break

        soup = BeautifulSoup(response.text, "html.parser")
        quotes = soup.find_all("div", class_="quote")

        for item in quotes:
            text_el = item.find("span", class_="text")
            author_el = item.find("small", class_="author")

            cleaned_text = (
                text_el.text.strip().strip('"').strip("“").strip("”")
                if text_el
                else "N/A"
            )
            cleaned_author = author_el.text.strip() if author_el else "N/A"

            # ဘာသာပြန်ခြင်း
            # ဘာသာပြန်ခြင်း (Error ခံထားခြင်း နှင့် 0.5 စက္ကန့် စောင့်ခိုင်းခြင်း)
            translated_mm = "N/A"
            if cleaned_text != "N/A":
                try:
                    translated_mm = translator.translate(cleaned_text)
                    time.sleep(0.5)  # Google Server မပိတ်စေရန် 0.5 စက္ကန့် စောင့်ခိုင်းခြင်း
                except Exception:
                    translated_mm = "(Translation Error)"

            tags_elements = item.find_all("a", class_="tag")
            tags_list = [t.text.strip() for t in tags_elements if t and t.text]
            cleaned_tags = ", ".join(tags_list)

            # Excel ထဲသို့ Row အဖြစ် ထည့်သွင်းခြင်း
            ws.append(
                [
                    quote_id,
                    cleaned_text,
                    translated_mm,
                    cleaned_author,
                    cleaned_tags,
                ]
            )
            quote_id += 1

        # Next Button (နောက်စာမျက်နှာ) ရှိမရှိ စစ်ဆေးခြင်း
        next_btn = soup.find("li", class_="next")
        if next_btn and next_btn.find("a"):
            next_href = next_btn.find("a")["href"]
            current_url = base_url + next_href  # URL အသစ်သို့ ပြောင်းခြင်း
            page_num += 1
        else:
            print("No more pages found. Scraping finished!")
            current_url = None  # Loop ကို ရပ်ရန် None သတ်မှတ်ခြင်း

    # Excel Styling ပြုလုပ်ခြင်း
    header_fill = PatternFill(
        start_color="1F4E78", end_color="1F4E78", fill_type="solid"
    )
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Pyidaungsu", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
    )

    # Header Style ထည့်ခြင်း
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # Data Style ထည့်ခြင်း
    for row in ws.iter_rows(
        min_row=2, max_row=ws.max_row, min_col=1, max_col=5
    ):
        for cell in row:
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = center_align if cell.column == 1 else left_align

    # Column Widths သတ်မှတ်ခြင်း
    column_widths = {"A": 8, "B": 45, "C": 45, "D": 20, "E": 25}
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    # File သိမ်းဆည်းခြင်း
    script_dir = os.path.dirname(os.path.abspath(__file__))
    output_file = os.path.join(script_dir, "all_pages_quotes_styled.xlsx")
    wb.save(output_file)

    print(f"Successfully saved all pages to: {output_file}")

except Exception as e:
    print(f"Error occurred: {e}")