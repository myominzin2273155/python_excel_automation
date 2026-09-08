import openpyxl

wb = openpyxl.load_workbook("styled_sales.xlsx")
sheet = wb["SalesData"]

data = []
for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, values_only=True):
    data.append(row)

sorted_data = sorted(data, key=lambda x: x[3], reverse=True)

sorted_sheet = wb.create_sheet(title="SortedData")

header = [cell.value for cell in sheet[1]]
sorted_sheet.append(header)

for row in sorted_data:
    sorted_sheet.append(row)

    sorted_sheet.auto_filter.ref = f"A1:D{sorted_sheet.max_row}"

    wb.save("filtered_sorted_sales.xlsx")

    print("Sorting , AutoFilter finish")
