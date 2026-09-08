import os
import openpyxl

master_wb = openpyxl.Workbook()
master_sheet = master_wb.active
master_sheet.title = "FolderMergedData"

master_sheet.append(["Item", "Total", "Source_File"])

current_directory = os.getcwd()

for file_name in os.listdir(current_directory):
    if file_name.endswith(".xlsx") and not file_name.startswith("master_") and not file_name.startswith("~$"):

        wb = openpyxl.load_workbook(file_name)
        sheet = wb.active

        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, values_only=True):

            row_data = list(row)
            row_data.append(file_name)
            master_sheet.append(row_data)

master_wb.save("master_all_folder_sales.xlsx")
print("Folder xJ&Sd Excel Merging finish")            