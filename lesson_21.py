import os
import time
import requests
from bs4 import BeautifulSoup
from deep_translator import GoogleTranslator
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ---------------------------------------------------------
# ၁။ ဘာသာပြန်ပေးသည့် Function (Translation Helper)
# ---------------------------------------------------------
def translate_text(translator, text):
    """စာသားကို မြန်မာလို ဘာသာပြန်ပေးပြီး Error ခံပေးသော Function"""
    if not text or text == "N/A":
        return "N/A"
    
    try:
        translated = translator.translate(text)
        time.sleep(0.5)  # Google Block မခံရစေရန် နားပေးခြင်း
        return translated
    except Exception:
        return "(Translation Error)"

# ---------------------------------------------------------
# ၂။ Excel File Formatting ပြုလုပ်ပေးသည့် Function
# ---------------------------------------------------------
def apply_excel_styles(ws):
    """Excel Sheet ကို အလှဆင်ပေးသော Function"""
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Pyidaungsu", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3")
    )

    # Header Styling
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # Data Rows Styling
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=5):
        for cell in row:
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = center_align if cell.column == 1 else left_align

    # Column Widths
    column_widths = {"A": 8, "B": 45, "C": 45, "D": 20, "E": 25}
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

# ---------------------------------------------------------
# ၃။ Main Automation Execution (အဓိက အလုပ်လုပ်သည့်နေရာ)
# ---------------------------------------------------------
def main():
    translator = GoogleTranslator(source="auto", target="my")
    wb = Workbook()
    ws = wb.active
    ws.title = "Refactored Quotes"
    ws.append(["ID", "Quote Text (Eng)", "Quote Text (MM)", "Author", "Tags"])

    base_url = "https://quotes.toscrape.com"
    current_url = base_url
    quote_id = 1
    page_num = 1

    try:
        while current_url:
            print(f"Fetching Page {page_num}: {current_url} ...")
            response = requests.get(current_url, timeout=10)
            if response.status_code != 200:
                print(f"Failed to fetch page {page_num}")
                break

            soup = BeautifulSoup(response.text, "html.parser")
            quotes = soup.find_all("div", class_="quote")

            for item in quotes:
                text_el = item.find("span", class_="text")
                author_el = item.find("small", class_="author")

                cleaned_text = text_el.text.strip().strip('"').strip('“').strip('”') if text_el else "N/A"
                cleaned_author = author_el.text.strip() if author_el else "N/A"

                # ဘာသာပြန် Function ကို သီးသန့် ခေါ်သုံးခြင်း
                translated_mm = translate_text(translator, cleaned_text)

                tags_elements = item.find_all("a", class_="tag")
                tags_list = [t.text.strip() for t in tags_elements if t and t.text]
                cleaned_tags = ", ".join(tags_list)

                ws.append([quote_id, cleaned_text, translated_mm, cleaned_author, cleaned_tags])
                quote_id += 1

            # Next Page Check
            next_btn = soup.find("li", class_="next")
            if next_btn and next_btn.find("a"):
                current_url = base_url + next_btn.find("a")["href"]
                page_num += 1
            else:
                print("Scraping completed!")
                current_url = None

        # Style Function ကို ခေါ်သုံးခြင်း
        apply_excel_styles(ws)

        # File သိမ်းခြင်း
        script_dir = os.path.dirname(os.path.abspath(__file__))
        output_file = os.path.join(script_dir, "lesson_21_refactored_quotes.xlsx")
        wb.save(output_file)
        print(f"Successfully saved to: {output_file}")

    except Exception as e:
        print(f"Error occurred: {e}")

# Program စတင်ပတ်သည့်နေရာ
if __name__ == "__main__":
    main()