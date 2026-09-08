import openpyxl

wb = openpyxl.load_workbook("master_all_folder_sales.xlsx")
sheet = wb.active

new_wb = openpyxl.Workbook()
new_sheet = new_wb.active
new_sheet.title = "CleanedData"

header = [cell.value for cell in sheet[1]]
new_sheet.append(header)

seen_rows = set()

for row in sheet.iter_rows(min_row=2, values_only=True):

    if not any(row):
        continue

    if  row not in seen_rows:
        seen_rows.add(row)
        new_sheet.append(row)

new_wb.save("cleaned_sales_data.xlsx")

print("Data Cleaning and Duplicate ဖျက်ခြင်း အောင်မြင့်စွာ ပြီးပါပြီ")