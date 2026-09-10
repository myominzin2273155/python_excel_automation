import os
from bs4 import BeautifulSoup
from deep_translator import GoogleTranslator
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import requests

translator = GoogleTranslator(source="auto", target="my")

wb = Workbook()
ws = wb.active
ws.title = "Styled Quotes"

headers = ["ID", "Quote Text (Eng)", "Quote Text (MM)", "Author", "Tags"]
ws.append(headers)

# ပြင်ဆင်ချက် ၁: wrl မှ url သို့ ပြင်ထားပါသည်
url = "https://quotes.toscrape.com/"

try:
    print("Fetching and translating data...")
    response = requests.get(url, timeout=10)

    if response.status_code == 200:
        soup = BeautifulSoup(response.text, "html.parser")
        quotes = soup.find_all("div", class_="quote")

        for index, item in enumerate(quotes, 1):
            text_el = item.find("span", class_="text")
            author_el = item.find("small", class_="author")

            cleaned_text = (
                text_el.text.strip().strip('"').strip("“").strip("”")
                if text_el
                else "N/A"
            )
            cleaned_author = author_el.text.strip() if author_el else "N/A"

            translated_mm = (
                translator.translate(cleaned_text)
                if cleaned_text != "N/A"
                else "N/A"
            )

            tags_elements = item.find_all("a", class_="tag")
            tags_list = [t.text.strip() for t in tags_elements if t and t.text]
            cleaned_tags = ", ".join(tags_list)

            ws.append(
                [
                    index,
                    cleaned_text,
                    translated_mm,
                    cleaned_author,
                    cleaned_tags,
                ]
            )

        # Styling
        header_fill = PatternFill(
            start_color="1F4E78", end_color="1F4E78", fill_type="solid"
        )
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        data_font = Font(name="Pyidaungsu", size=11)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(
            horizontal="left", vertical="center", wrap_text=True
        )

        # ပြင်ဆင်ချက် ၂: "D3D3D3" တွင် quote ထည့်ထားပါသည်
        thin_border = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3"),
        )

        # Apply Header Style
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        # Apply Data Style
        for row in ws.iter_rows(
            min_row=2, max_row=ws.max_row, min_col=1, max_col=5
        ):
            for cell in row:
                cell.font = data_font
                cell.border = thin_border

                if cell.column == 1:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

        # ပြင်ဆင်ချက် ၃: Column Widths နှင့် Save အပိုင်းကို Loop အပြင်ဘက်သို့ ထုတ်ထားပါသည်
        column_widths = {"A": 8, "B": 45, "C": 45, "D": 20, "E": 25}
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        script_dir = os.path.dirname(os.path.abspath(__file__))
        output_file = os.path.join(script_dir, "cleaned_quotes_styled.xlsx")
        wb.save(output_file)

        print(f"Success! Output saved to: {output_file}")

    else:
        print(f"Connection Failed: {response.status_code}")

except Exception as e:
    print(f"Error occurred: {e}")